from __future__ import annotations

import hashlib
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from .common import ActivityError, digest

SHEET_NAME = "物料进出库明细报表"
HEADERS = ["序号", "物料编码", "物料名称", "物料规格", "物料类别", "交易日期", "交易类型",
           "单位", "数量", "单价（元）", "金额（元）", "出入库部门", "经办人", "供应商名称", "备注"]
REQUIRED_HEADERS = {"物料编码", "物料名称", "交易日期", "交易类型", "数量", "单位"}


def raw_cell(cell: Any) -> dict:
    value = cell.value
    if isinstance(value, (date, datetime)):
        value = value.isoformat()
    elif isinstance(value, float) and (value != value or abs(value) == float("inf")):
        value = str(value)
    return {"value": value, "type": cell.data_type, "number_format": cell.number_format}


def load_erp(content: bytes, filename: str) -> tuple[list[dict], dict]:
    """Read raw cells (including formulas), never replace them with normalized values."""
    file_hash = hashlib.sha256(content).hexdigest()
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ActivityError("无法读取 Excel 工作簿") from exc
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ActivityError(f"缺少工作表: {SHEET_NAME}")
        sheet = workbook[SHEET_NAME]
        first = next(sheet.iter_rows(min_row=1, max_row=1), ())
        nonempty = [(i, str(c.value).strip()) for i, c in enumerate(first) if c.value is not None]
        headers = [h for _, h in nonempty]
        if len(headers) != len(set(headers)):
            raise ActivityError("表头重复，无法可靠映射字段")
        missing = REQUIRED_HEADERS - set(headers)
        if missing:
            raise ActivityError("缺少关键字段: " + "、".join(sorted(missing)))
        # Some ERP workbooks have formatting extending to XFD. Restrict to real headers.
        last_column = max(i for i, _ in nonempty) + 1
        records = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=last_column), 2):
            if all(c.value is None or str(c.value).strip() == "" for c in cells):
                continue
            records.append({
                "source_record_id": "source/" + digest([file_hash, SHEET_NAME, row_number]),
                "source_file_sha256": file_hash, "source_filename": filename,
                "sheet_name": SHEET_NAME, "excel_row": row_number,
                "raw_values": {h: raw_cell(cells[i]) for i, h in nonempty},
                "normalized_values": {}, "issues": [], "status": "pending",
            })
        if not records:
            raise ActivityError("工作表没有非空数据记录")
        return records, {"source_file_sha256": file_hash, "source_filename": filename,
                         "sheet_name": SHEET_NAME, "headers": headers,
                         "excel_epoch": workbook.epoch.isoformat()}
    finally:
        workbook.close()
