from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .models import DataQualityError, Transaction


REQUIRED_FIELDS = {
    "material_code",
    "material_name",
    "material_spec",
    "trans_date",
    "trans_type",
    "quantity",
    "amount",
}

FIELD_ALIASES: dict[str, list[str]] = {
    "material_code": ["物料编码", "物料代码", "物料编号", "MaterialCode", "产品编码"],
    "material_name": ["物料名称", "物料名", "品名", "MaterialName", "material_desc", "产品名称"],
    "material_spec": ["物料规格", "规格", "型号", "规格型号", "Spec", "产品规格"],
    "trans_date": ["交易日期", "日期", "业务日期", "出入库日期", "Date"],
    "trans_type": ["交易类型", "出入库类型", "类型", "收发标志", "Type"],
    "quantity": ["数量", "出入库数量", "收发数量", "Qty", "件数"],
    "amount": ["金额", "金额（元）", "单价金额", "出入库金额", "Amt"],
    "material_category": ["物料类别", "物料分类", "品类", "Category", "产品类别"],
    "dept_name": ["出入库部门", "部门", "仓库", "Dept"],
    "operator": ["经办人", "制单人", "操作员", "Operator"],
    "supplier": ["供应商名称", "供应商", "供货单位", "Supplier"],
}

TYPE_IN = {"入库", "in", "1", "收", "采购入库", "生产入库", "外协入库"}
TYPE_OUT = {"出库", "out", "-1", "发", "销售出库", "生产领料", "委外出库"}


def parse_workbook(content: bytes) -> tuple[list[Transaction], dict[str, Any]]:
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    sheet, mapping, headers = _select_sheet(workbook)

    transactions: list[Transaction] = []
    invalid_rows: list[str] = []
    total_rows = 0
    recognized_types = 0

    for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if _is_empty_row(row):
            continue
        total_rows += 1
        try:
            tx = _parse_row(row, mapping)
            recognized_types += 1
            transactions.append(tx)
        except DataQualityError as exc:
            invalid_rows.append(f"第{row_index}行: {exc.message}")

    data_quality = _validate_quality(transactions, total_rows, invalid_rows, recognized_types)
    data_quality.update(
        {
            "sheet_name": sheet.title,
            "headers": headers,
            "field_mapping": {field: headers[index] for field, index in mapping.items() if index < len(headers)},
            "invalid_samples": invalid_rows[:10],
        }
    )
    return transactions, data_quality


def _select_sheet(workbook: Any) -> tuple[Any, dict[str, int], list[str]]:
    candidates: list[tuple[int, Any, dict[str, int], list[str]]] = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(min_row=1, max_row=1, values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        mapping = _map_headers(headers)
        score = len(REQUIRED_FIELDS.intersection(mapping))
        if REQUIRED_FIELDS.issubset(mapping):
            candidates.append((score, sheet, mapping, headers))

    if not candidates:
        first_headers = []
        if workbook.worksheets:
            rows = workbook.worksheets[0].iter_rows(min_row=1, max_row=1, values_only=True)
            first_headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        missing = "、".join(sorted(REQUIRED_FIELDS))
        raise DataQualityError(
            "数据格式不符合标准",
            [
                f"必须字段缺失，无法找到同时包含这些字段的 sheet: {missing}",
                f"当前第一个 sheet 字段: {', '.join(first_headers)}",
            ],
        )

    candidates.sort(key=lambda item: item[0], reverse=True)
    _, sheet, mapping, headers = candidates[0]
    return sheet, mapping, headers


def _map_headers(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    normalized_headers = [_normalize_header(header) for header in headers]
    for field, aliases in FIELD_ALIASES.items():
        alias_set = {_normalize_header(alias) for alias in aliases}
        for index, normalized in enumerate(normalized_headers):
            if normalized in alias_set:
                mapping[field] = index
                break
        if field in mapping:
            continue
        for index, normalized in enumerate(normalized_headers):
            if any(alias and alias in normalized for alias in alias_set):
                mapping[field] = index
                break
    return mapping


def _parse_row(row: tuple[Any, ...], mapping: dict[str, int]) -> Transaction:
    code = _text(_value(row, mapping["material_code"]))
    if not code or not re.match(r"^[A-Za-z0-9_-]+$", code):
        raise DataQualityError("物料编码为空或格式错误")

    name = _text(_value(row, mapping["material_name"]))
    if not name:
        raise DataQualityError("物料名称为空")

    trans_type = _normalize_type(_value(row, mapping["trans_type"]))
    if trans_type is None:
        raise DataQualityError("交易类型无法识别")

    return Transaction(
        material_code=code,
        material_name=name,
        material_spec=_text(_value(row, mapping["material_spec"])),
        trans_date=_parse_date(_value(row, mapping["trans_date"])),
        trans_type=trans_type,
        quantity=_parse_number(_value(row, mapping["quantity"]), "数量"),
        amount=_parse_number(_value(row, mapping["amount"]), "金额"),
        material_category=_optional_text(row, mapping, "material_category"),
        dept_name=_optional_text(row, mapping, "dept_name"),
        operator=_optional_text(row, mapping, "operator"),
        supplier=_optional_text(row, mapping, "supplier"),
    )


def _validate_quality(
    transactions: list[Transaction],
    total_rows: int,
    invalid_rows: list[str],
    recognized_types: int,
) -> dict[str, Any]:
    if total_rows == 0:
        raise DataQualityError("数据为空", ["库存明细 sheet 没有可分析的数据行"])

    invalid_ratio = len(invalid_rows) / total_rows
    if invalid_ratio >= 0.1:
        raise DataQualityError(
            "数据质量过低",
            [f"无效行比例: {invalid_ratio:.1%}", "必须字段空值或格式错误行比例不能超过 10%"] + invalid_rows[:10],
        )

    type_ratio = recognized_types / total_rows
    if type_ratio <= 0.8:
        raise DataQualityError("交易类型不完整", [f"可识别交易类型占比: {type_ratio:.1%}", "最低要求: >80%"])

    dates = [tx.trans_date for tx in transactions]
    if (max(dates) - min(dates)).days < 30:
        raise DataQualityError(
            "数据周期不足",
            [f"数据覆盖天数: {(max(dates) - min(dates)).days}天", "最低要求: 30天"],
        )

    material_count = len({tx.material_code for tx in transactions})
    if material_count < 5:
        raise DataQualityError("物料数量不足以分析", [f"不同物料数: {material_count}", "最低要求: 5种"])

    total_amount = sum(tx.amount for tx in transactions)
    if total_amount <= 0:
        raise DataQualityError("金额数据异常", ["总金额必须大于 0"])

    return {
        "total_records": total_rows,
        "valid_records": len(transactions),
        "invalid_records": len(invalid_rows),
        "quality_grade": _quality_grade(invalid_ratio),
        "date_range": {"start_date": min(dates).isoformat(), "end_date": max(dates).isoformat()},
        "material_count": material_count,
    }


def _quality_grade(invalid_ratio: float) -> str:
    if invalid_ratio == 0:
        return "A"
    if invalid_ratio < 0.03:
        return "B"
    return "C"


def _normalize_type(value: Any) -> str | None:
    normalized = _text(value).lower()
    if normalized in TYPE_IN:
        return "入库"
    if normalized in TYPE_OUT:
        return "出库"
    return None


def _parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return from_excel(value).date()

    text = _text(value)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise DataQualityError("日期格式无法识别")


def _parse_number(value: Any, label: str) -> float:
    if value is None or value == "":
        raise DataQualityError(f"{label}为空")
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace(",", "")
        try:
            number = float(text)
        except ValueError as exc:
            raise DataQualityError(f"{label}不是数字") from exc
    if number < 0:
        raise DataQualityError(f"{label}不能为负数")
    return number


def _optional_text(row: tuple[Any, ...], mapping: dict[str, int], field: str) -> str | None:
    if field not in mapping:
        return None
    value = _text(_value(row, mapping[field]))
    return value or None


def _value(row: tuple[Any, ...], index: int) -> Any:
    if index >= len(row):
        return None
    return row[index]


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_header(value: str) -> str:
    return re.sub(r"\s+", "", value.strip().lower())


def _is_empty_row(row: tuple[Any, ...]) -> bool:
    return all(value is None or str(value).strip() == "" for value in row)
